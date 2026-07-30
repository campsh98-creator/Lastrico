from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FINAL_DIR = ROOT / "docs" / "business-plan" / "final"
ASSET_DIR = FINAL_DIR / "assets"
DOCX_PATH = FINAL_DIR / "MOTIA_BUSINESS_PLAN_2026_2029.docx"
MODEL_PATH = ROOT / "financials" / "MOTIA_FINANCIAL_MODEL.csv"
FUNDING_PATH = ROOT / "financials" / "MOTIA_FUNDING_SCENARIOS.csv"

# Narrative-proposal token map, explicitly adapted to A4 and the provisional Motia system.
THEME = {
    "font_heading": "Arial",
    "font_body": "Arial",
    "midnight": "132A2A",
    "teal": "0D8B73",
    "coral": "F36C4F",
    "paper": "F6F1E7",
    "mist": "DDE7E3",
    "slate": "344443",
    "light": "F4F6F5",
    "white": "FFFFFF",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    bold_candidates = [
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
    ]
    for candidate in bold_candidates if bold else candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                pass
    return ImageFont.load_default()


def read_model() -> tuple[list[dict[str, str]], dict[str, str]]:
    with MODEL_PATH.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    base = [row for row in rows if row["Scenario"] == "Base"]
    with FUNDING_PATH.open(newline="", encoding="utf-8") as handle:
        funding = next(row for row in csv.DictReader(handle) if row["Scenario"] == "Base")
    return base, funding


def n(row: dict[str, str], key: str) -> float:
    value = row.get(key, "")
    return float(value) if value not in ("", None) else 0.0


def euro(value: float, decimals: int = 0) -> str:
    rendered = f"{value:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"€{rendered}"


def compact_euro(value: float) -> str:
    if abs(value) >= 1_000_000:
        return f"€{value / 1_000_000:.1f}m".replace(".", ",")
    if abs(value) >= 1_000:
        return f"€{value / 1_000:.0f}k"
    return euro(value)


def create_brand_mark(path: Path) -> None:
    image = Image.new("RGBA", (900, 260), "#F6F1E7")
    draw = ImageDraw.Draw(image)
    points = [(55, 198), (55, 86), (122, 158), (180, 76), (244, 187), (270, 187)]
    draw.line(points, fill="#0D8B73", width=28, joint="curve")
    for point in (points[0], points[-1]):
        x, y = point
        draw.ellipse((x - 14, y - 14, x + 14, y + 14), fill="#0D8B73")
    draw.text((320, 69), "MOTIA", font=load_font(92, True), fill="#132A2A")
    draw.text((325, 165), "Routes shaped around you.", font=load_font(27), fill="#344443")
    image.save(path, dpi=(300, 300))


def chart_canvas(title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), title, font=load_font(46, True), fill="#132A2A")
    if subtitle:
        draw.text((90, 115), subtitle, font=load_font(24), fill="#344443")
    return image, draw


def draw_line_chart(
    path: Path,
    title: str,
    labels: list[str],
    series: list[tuple[str, list[float], str]],
    money: bool = False,
) -> None:
    image, draw = chart_canvas(title, "Scenario Base · ipotesi gestionali, non previsione")
    left, top, right, bottom = 150, 200, 1690, 735
    all_values = [value for _, values, _ in series for value in values]
    maximum = max(all_values) if all_values else 1
    maximum = maximum * 1.12 if maximum else 1
    for step in range(5):
        y = bottom - (bottom - top) * step / 4
        draw.line((left, y, right, y), fill="#DDE7E3", width=2)
        value = maximum * step / 4
        label = compact_euro(value) if money else (f"{value / 1000:.0f}k" if value >= 1000 else f"{value:.0f}")
        draw.text((30, y - 14), label, font=load_font(22), fill="#344443")
    draw.line((left, top, left, bottom), fill="#344443", width=3)
    draw.line((left, bottom, right, bottom), fill="#344443", width=3)
    x_positions = [
        left + (right - left) * index / max(1, len(labels) - 1) for index in range(len(labels))
    ]
    for index, label in enumerate(labels):
        if index % 2 == 0 or index == len(labels) - 1:
            draw.text((x_positions[index] - 38, bottom + 22), label, font=load_font(18), fill="#344443")
    for name, values, colour in series:
        points = []
        for x, value in zip(x_positions, values):
            y = bottom - (bottom - top) * value / maximum
            points.append((x, y))
        # A plain polyline avoids PIL's rounded-joint artefact on sharp financing spikes.
        draw.line(points, fill=colour, width=7)
        for x, y in points:
            draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill=colour)
    legend_x = left
    for name, _, colour in series:
        draw.line((legend_x, 830, legend_x + 48, 830), fill=colour, width=7)
        draw.text((legend_x + 60, 813), name, font=load_font(22), fill="#344443")
        legend_x += 420
    image.save(path, dpi=(300, 300))


def draw_use_of_funds(path: Path) -> None:
    categories = [
        ("Prodotto, routing e mobile", 38, "#0D8B73"),
        ("Dati, moderazione e infrastruttura", 23, "#277F87"),
        ("Legale, privacy, sicurezza e assicurazione", 10, "#627D79"),
        ("Piloti, field test e commerciale", 17, "#F36C4F"),
        ("Runway fondatore/core team e riserva", 12, "#A6B4B0"),
    ]
    image, draw = chart_canvas("Uso indicativo del pre-seed Base", "Allocazione preliminare del netto iniziale · validazione richiesta")
    left, top, right = 620, 220, 1650
    for index, (label, value, colour) in enumerate(categories):
        y = top + index * 110
        draw.text((90, y + 14), label, font=load_font(25, True if index == 0 else False), fill="#132A2A")
        draw.rounded_rectangle((left, y, right, y + 56), radius=18, fill="#EDF2F0")
        width = (right - left) * value / 45
        draw.rounded_rectangle((left, y, left + width, y + 56), radius=18, fill=colour)
        draw.text((right + 25, y + 9), f"{value}%", font=load_font(28, True), fill="#132A2A")
    draw.text((90, 800), "Non include automaticamente il follow-on 2028: ogni tranche resta milestone-gated.", font=load_font(23), fill="#344443")
    image.save(path, dpi=(300, 300))


def draw_architecture(path: Path) -> None:
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 45), "Architettura del valore Motia", font=load_font(46, True), fill="#132A2A")
    boxes = [
        (90, 210, 430, 500, "INTENTO", "Prompt o controlli\nespliciti dell'utente", "#F6F1E7"),
        (525, 210, 865, 500, "POLICY", "Schema bounded,\nconfermato e versionato", "#DDE7E3"),
        (960, 210, 1300, 500, "ROUTING", "Candidati legali,\nscore deterministico", "#DDE7E3"),
        (1395, 210, 1735, 500, "ESPERIENZA", "Percorso spiegato,\nlimiti e incertezza", "#F6F1E7"),
    ]
    for index, (x1, y1, x2, y2, heading, body, colour) in enumerate(boxes):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=colour, outline="#0D8B73", width=4)
        draw.text((x1 + 30, y1 + 40), heading, font=load_font(30, True), fill="#0D8B73")
        draw.multiline_text((x1 + 30, y1 + 120), body, font=load_font(28), fill="#132A2A", spacing=18)
        if index < len(boxes) - 1:
            draw.line((x2 + 15, 355, boxes[index + 1][0] - 15, 355), fill="#F36C4F", width=8)
    draw.text((90, 615), "L'AI interpreta preferenze supportate. Legalità, grafo, scoring e ricalcolo restano deterministici.", font=load_font(28, True), fill="#132A2A")
    image.save(path, dpi=(300, 300))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = THEME["font_body"]
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = rgb(THEME["slate"])
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, colour in [
        ("Title", 30, THEME["midnight"]),
        ("Heading 1", 20, THEME["midnight"]),
        ("Heading 2", 13.5, THEME["teal"]),
        ("Heading 3", 11, THEME["midnight"]),
    ]:
        style = styles[name]
        style.font.name = THEME["font_heading"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(colour)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.text = "MOTIA   ·   Investor Working Draft   ·   July 2026"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.name = THEME["font_body"]
    p.runs[0].font.size = Pt(7.5)
    p.runs[0].font.color.rgb = rgb("627D79")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Motia · Investor Working Draft · ")
    run.font.name = THEME["font_body"]
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb("627D79")
    add_page_field(p)


def add_section_label(document: Document, number: str, title: str, kicker: str | None = None) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(number.upper())
    run.bold = True
    run.font.name = THEME["font_heading"]
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(THEME["coral"])
    document.add_heading(title, level=1)
    if kicker:
        p = document.add_paragraph(kicker)
        p.style = document.styles["Subtitle"]
        p.runs[0].font.name = THEME["font_body"]
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.color.rgb = rgb(THEME["teal"])


def add_body(document: Document, text: str, bold_lead: str | None = None) -> None:
    p = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first, rest = text.split(":", 1)
        run = p.add_run(first + ":")
        run.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(document: Document, text: str, fill: str = "DDE7E3") -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 160, 220, 160, 220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.name = THEME["font_heading"]
    run.font.size = Pt(14)
    run.font.color.rgb = rgb(THEME["midnight"])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, THEME["midnight"])
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        run = p.add_run(heading)
        run.bold = True
        run.font.size = Pt(8.3)
        run.font.color.rgb = rgb(THEME["white"])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, THEME["light"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(8.1)
            run.font.color.rgb = rgb(THEME["slate"])
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(document: Document, path: Path, width_cm: float, caption: str | None = None) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        p = document.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(7.5)
            run.font.italic = True
            run.font.color.rgb = rgb("627D79")


def page_break(document: Document) -> None:
    p = document.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def sum_year(rows: list[dict[str, str]], key: str, year: str) -> float:
    return sum(n(row, key) for row in rows if year in row["Quarter"])


def build() -> None:
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    base, funding = read_model()

    logo_path = ASSET_DIR / "motia-wordmark-generated.png"
    users_chart = ASSET_DIR / "base-user-growth.png"
    finance_chart = ASSET_DIR / "base-financial-trajectory.png"
    use_funds_chart = ASSET_DIR / "base-use-of-funds.png"
    architecture_chart = ASSET_DIR / "motia-architecture.png"
    create_brand_mark(logo_path)
    labels = [row["Quarter"].replace(" ", "\n") for row in base]
    draw_line_chart(
        users_chart,
        "Crescita utenti attivi mensili",
        labels,
        [("MAU", [n(row, "MAU") for row in base], "#0D8B73")],
    )
    draw_line_chart(
        finance_chart,
        "Traiettoria economica e cassa",
        labels,
        [
            ("Ricavi operativi", [n(row, "Total_Operating_Revenue_EUR") for row in base], "#0D8B73"),
            ("Costi operativi", [n(row, "Total_Operating_Costs_EUR") for row in base], "#F36C4F"),
            ("Cassa finale", [n(row, "Closing_Cash_EUR") for row in base], "#132A2A"),
        ],
        money=True,
    )
    draw_use_of_funds(use_funds_chart)
    draw_architecture(architecture_chart)

    document = Document()
    configure_document(document)
    core = document.core_properties
    core.title = "MOTIA — Business Plan 2026–2029"
    core.subject = "Investor Working Draft"
    core.author = "Domenico Campanella Scali — founder verification required"
    core.keywords = "Motia, Lastrico, routing intelligence, business plan, mobility"
    core.comments = "Version 1.0 — July 2026"

    # 1 — Cover
    document.add_paragraph().paragraph_format.space_after = Pt(18)
    add_picture(document, logo_path, 15.5)
    p = document.add_paragraph("BUSINESS PLAN")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.name = THEME["font_heading"]
    p.runs[0].font.size = Pt(28)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = rgb(THEME["midnight"])
    p = document.add_paragraph("2026–2029")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(22)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = rgb(THEME["teal"])
    document.add_paragraph().paragraph_format.space_after = Pt(18)
    add_callout(document, "Lastrico by Motia  /  Smoother roads. Smarter routes.", THEME["paper"])
    p = document.add_paragraph("Investor Working Draft\nVersion 1.0 · July 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = rgb(THEME["slate"])
    p = document.add_paragraph("\nFounder: Domenico Campanella Scali")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = rgb(THEME["midnight"])
    p = document.add_paragraph("Nome e profilo biografico da verificare prima della pubblicazione esterna.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.italic = True
    p.runs[0].font.color.rgb = rgb("627D79")
    page_break(document)

    # 2 — Executive summary
    add_section_label(document, "01", "Executive summary", "Un programma di validazione, non un lancio nazionale")
    add_body(document, "Motia è una proposta di startup di routing intelligence: vuole permettere a persone e organizzazioni di scegliere percorsi coerenti con veicolo, strada, incertezza e preferenze reali, non solo con il tempo stimato.")
    add_callout(document, "La futura interfaccia della navigazione è il prompt.")
    add_body(document, "La frase non affida la navigazione a un modello linguistico. Il prompt viene tradotto in uno schema limitato e confermabile; legalità, generazione dei candidati, scoring e ricalcolo restano deterministici e verificabili.")
    add_table(
        document,
        ["Livello", "Ruolo", "Stato"],
        [
            ["Motia", "Società madre · Routes shaped around you.", "Proposta; verifiche societarie e marchio"],
            ["Motia Routing Intelligence", "Preferenze, condizioni, ranking, spiegazioni", "Piattaforma pianificata"],
            ["Motia Routing Engine", "Servizio/API per partner", "Visione futura"],
            ["Lastrico by Motia", "Primo prodotto surface-aware", "Beta sperimentale testata a Milano"],
        ],
        [4.2, 8.0, 4.2],
    )
    add_body(document, "La tesi investibile è misurata: finanziare quattro prove — comportamento ripetuto, qualità dati, delivery nativa fedele al percorso selezionato ed economia sostenibile — prima di estendere piattaforme, geografie o superfici automotive.")
    page_break(document)

    # 3 — Vision, mission, problem
    add_section_label(document, "02", "Visione, missione e problema")
    document.add_heading("Missione", level=2)
    add_body(document, "Permettere a persone e organizzazioni di scegliere percorsi realmente adatti alle proprie esigenze, traducendo preferenze complesse in decisioni di navigazione intelligenti, trasparenti e personalizzabili.")
    document.add_heading("Visione", level=2)
    add_body(document, "Costruire un nuovo livello di intelligenza per la mobilità, nel quale ogni percorso possa essere modellato attraverso il linguaggio naturale e ottimizzato sulla base di ciò che conta davvero per l'utente.")
    document.add_heading("Il problema", level=2)
    add_bullets(document, [
        "Le interfacce mainstream mettono spesso in primo piano l'ETA; una scorciatoia può essere inadatta al mezzo o alle preferenze.",
        "Qualità del fondo, restringimenti, pendenze, continuità della viabilità principale, chiusure e affidabilità hanno significati diversi per auto, moto e bici.",
        "I dati sono frammentati, incompleti o non aggiornati; l'assenza di segnalazioni non equivale a una strada sicura.",
        "Gli utenti non dispongono di un modo semplice per esprimere compromessi complessi e verificabili.",
    ])
    add_callout(document, "173.364 incidenti con lesioni · 3.030 vittime · €22,6 mld di costo sociale stimato in Italia nel 2024", THEME["paper"])
    add_body(document, "Fonte: Istat, Incidenti stradali in Italia 2024. Questi dati descrivono la rilevanza del dominio; non dimostrano che Motia riduca gli incidenti.")
    page_break(document)

    # 4 — Lastrico
    add_section_label(document, "03", "Lastrico: il primo prodotto", "Smoother roads. Smarter routes.")
    add_body(document, "Lastrico è l'evidenza tecnica iniziale: una PWA open source che confronta il percorso più rapido con un'alternativa a minore esposizione stimata a fondi irregolari noti, quando esiste un candidato realmente distinto.")
    add_table(
        document,
        ["Oggi verificabile", "Non ancora disponibile"],
        [
            ["React/Next.js PWA; MapLibre e dati OSM", "App native iOS/Android e navigazione background production-grade"],
            ["Profili auto, moto e bici", "CarPlay o Android Auto approvati"],
            ["Valhalla; fallback OSRM solo auto", "SLA, traffico live e fornitori commerciali contrattualizzati"],
            ["GPS foreground, manovre, voce opzionale", "Prompt routing e personalizzazione validata"],
            ["Report con stati pending/verified/rejected", "Moderazione operativa nazionale dimostrata"],
            ["Copertura beta testata a Milano", "Copertura verificata dell'Italia"],
        ],
        [8.2, 8.2],
    )
    add_body(document, "Milano è il primo banco di prova, non il nome del prodotto né un limite architetturale. L'estensione avverrà per zone e corridoi versionati, mostrando copertura nota e quota non valutata.")
    add_body(document, "Lastrico non garantisce un percorso sicuro o privo di pericoli. Segnaletica, leggi, condizioni attuali e giudizio dell'utente hanno sempre priorità.")
    page_break(document)

    # 5 — Product and technology
    add_section_label(document, "04", "Prodotto e tecnologia", "AI confinata, routing deterministico")
    add_picture(document, architecture_chart, 16.4)
    add_bullets(document, [
        "Schema di preferenze versionato, allow-listed e confermato dall'utente.",
        "PostgreSQL/PostGIS, ingestione versionata, identificatori stradali stabili e snapshot riproducibili.",
        "Contratto di selezione del percorso con candidate ID, geometria/token provider, versioni dati e motivo di reroute.",
        "Dati con fonte, licenza, tempo, freschezza, confidenza, rilevanza per modalità, stato ed expiry.",
        "App nativa primaria prima; seconda piattaforma e automotive solo dopo evidenza.",
    ])
    add_body(document, "Strategia provider: partire da una base di navigazione contrattualizzata, mantenendo in Motia il differenziale di dati, preferenze, ranking e spiegazione. Il self-hosting completo diventa sensato solo quando scala e controllo compensano il carico operativo.")
    page_break(document)

    # 6 — Market and competition
    add_section_label(document, "05", "Mercato, clienti e concorrenza")
    add_body(document, "Il piano non pubblica un TAM non verificato. La ricerca deve costruire un mercato bottom-up per segmento, corridoio, numero di flotte, budget, procurement e disponibilità a pagare.")
    add_table(
        document,
        ["Segmento iniziale", "Test 2026–27", "Segnale per continuare"],
        [
            ["Utenti surface-sensitive", "40+ interviste e route-choice cohort", "≥30% sceglie alternativa pertinente; target D30 ≥20%"],
            ["Flotte e assistenza", "5 design partner; 3 pilot scope", "3 pilot pagati o contrattualmente impegnati"],
            ["Comuni/operatori", "10 buyer interview; prototipo dati", "Budget owner, licenza/fonte e pilot pagato o cofinanziato"],
        ],
        [4.2, 6.0, 6.4],
    )
    document.add_heading("Posizionamento competitivo", level=2)
    add_body(document, "Motia non compete inizialmente sul basemap. Integra fondazioni mature e concentra la differenziazione su cinque elementi: rete di condizioni verificata, contratto di preferenze, evidenza di ranking, workflow per partner/contributori e fiducia attraverso limiti e incertezza visibili.")
    add_body(document, "Google, Waze, Apple, Mapbox e stack OSM rendono alta la pressione competitiva. Un'interfaccia conversazionale, da sola, è replicabile; il vantaggio deve emergere da dati, calibrazione, integrazione e operazioni.")
    page_break(document)

    # 7 — Community and marketing
    add_section_label(document, "06", "Community, open source e marketing")
    add_body(document, "L'open source è una strategia di fiducia e qualità, non crescita automatica. Client, test, documentazione e strumenti di valutazione possono restare pubblici; feed verificati, moderazione, SLA, integrazioni e API gestite possono sostenere servizi a pagamento.")
    add_bullets(document, [
        "Percorsi guidati per codice, dati, accessibilità, documentazione, localizzazione e test.",
        "Contributi non tecnici tramite report in-app, interviste, validation day e partner locali.",
        "Report non verificati non modificano immediatamente il routing; correzione, cancellazione e appello restano disponibili.",
        "Metriche: contributi verificati, rejection rate, tempo di moderazione, freschezza, retention contributori e km non noti.",
    ])
    document.add_heading("Go-to-market a budget contenuto", level=2)
    add_bullets(document, [
        "Aggiornamento settimanale del fondatore, basato su evidenze e limiti reali.",
        "Case study di percorsi, contenuti SEO, YouTube tecnico-divulgativo e newsletter proprietaria.",
        "Partnership con community locali; earned media solo su risultati misurati.",
        "Paid media limitato a esperimenti dopo aver compreso attivazione e messaggio organico.",
    ])
    add_callout(document, "Budget marketing dettagliato Base 2026–2029: €110.150. Il resto della voce commerciale finanzia vendite, partnership, delivery pilot e riserva.", THEME["paper"])
    page_break(document)

    # 8 — Business model
    add_section_label(document, "07", "Business model e unit economics")
    add_table(
        document,
        ["Linea", "Ipotesi di valore", "Gate"],
        [
            ["Free + founding member", "Confronto percorsi e contributi; supporto iniziale", "Attivazione e retention"],
            ["Premium B2C", "Preferenze, alert e spiegazioni avanzate", "Prezzo, churn, store effects, CAC/payback"],
            ["Fleet / B2B", "Route intelligence e pilot per operazioni", "ACV, deployment cost, rinnovo"],
            ["B2G / mapping", "Dashboard o progetto cofinanziato", "Procurement, licenze, budget owner"],
            ["API / white label", "Scoring e condition intelligence", "SLA, margine, diritti provider"],
            ["Grant / ricerca", "Finanza non diluitiva o income contrattuale", "Elegibilità e trattamento contabile"],
        ],
        [3.8, 7.0, 5.8],
    )
    add_body(document, "Ipotesi B2C: €4,99/mese o €39,99/anno di listino; ARPPU netto trimestrale Base €10. L'importo deve assorbire mix mensile/annuale, sconti, IVA, commissioni store, rimborsi, pagamenti falliti e churn.")
    add_bullets(document, [
        "Costo e margine per sessione di navigazione a 10k, 100k e 250k MAU.",
        "Conversione, ARPU netto, churn, CAC e payback per coorte.",
        "ACV, margine lordo, ciclo vendita, costo implementazione e rinnovo B2B/B2G.",
        "Costo di dati/moderazione per km verificato e regione attiva.",
    ])
    add_body(document, "Motia non venderà storici di movimento identificabili. La pubblicità non è prioritaria perché può compromettere fiducia, leggibilità, privacy e distrazione.")
    page_break(document)

    # 9 — Roadmap
    add_section_label(document, "08", "Roadmap 2026–2029", "Ogni fase è condizionale")
    add_table(
        document,
        ["Periodo", "Outcome target", "Gate principale"],
        [
            ["Q4 2026", "Super-beta controllata e baseline", "Stabilità, claims/privacy, 40 interviste"],
            ["H1 2027", "Backend regionale e alpha nativa primaria", "Route fidelity, moderazione, 3 pilot"],
            ["H2 2027", "Field beta e test prezzi", "D30, route acceptance, cost/session"],
            ["H1 2028", "Seconda regione e prompt-to-schema", "Data quality e corpus ambiguità"],
            ["H2 2028", "Motia Routing Intelligence beta", "B2B ripetibile, reliability, margine"],
            ["H1 2029", "API alpha; automotive solo se eligible", "Checklist piattaforme e domanda"],
            ["H2 2029", "Decisione multi-regione", "Economia corridoi, rinnovi, qualità auditata"],
        ],
        [2.7, 7.2, 6.7],
    )
    add_picture(document, users_chart, 15.7, "MAU Base: target gestionali scenario-based; non rappresentano una previsione.")
    add_body(document, "Stop o restringimento se gli utenti non comprendono/scelgono le alternative, se manca un buyer con valore misurabile, se il provider non preserva il percorso, se i dati non raggiungono qualità sostenibile o se il capitale richiede diluizione sproporzionata.")
    page_break(document)

    # 10 — Founder
    add_section_label(document, "09", "Fondatore e team", "Verifica fondatore richiesta prima della pubblicazione esterna")
    document.add_heading("Domenico Campanella Scali", level=2)
    add_body(document, "Profilo fornito dal fondatore: italiano; laurea in Economia Aziendale presso Università Politecnica delle Marche; laurea magistrale in Management e Strategia d'Impresa presso Università degli Studi di Verona; Master in Fintech in corso presso Politecnico di Milano Graduate School of Management.")
    add_body(document, "Il brief descrive inoltre esperienze in business analysis, servizi finanziari, progetti bancari, tecnologia, iniziative connesse all'AI, prototipi software e workflow AI-enabled. Ogni nome, titolo, stato e esperienza richiede conferma documentale.")
    add_callout(document, "Punto di forza: collegare business, finanza, tecnologia e problema utente, costruendo pubblicamente senza inventare competenze.", THEME["paper"])
    document.add_heading("Competenze da aggiungere", level=2)
    add_bullets(document, [
        "lead geospatial e routing;",
        "engineering iOS/Android e backend/data;",
        "product design, ricerca e accessibilità;",
        "community, moderazione e data operations;",
        "supporto legale/privacy/security, assicurativo, finanziario e fundraising.",
    ])
    add_body(document, "Il piano non presenta Domenico come senior AI researcher, software architect, geospatial scientist, serial entrepreneur o mobility executive. Governance, documentazione, advisor indipendenti e ownership distribuita mitigano la dipendenza dal singolo fondatore.")
    page_break(document)

    # 11 — Funding
    add_section_label(document, "10", "Funding ed equity crowdfunding")
    gross_target = n(funding, "Gross_Funding_Target_EUR")
    transaction_cost = n(funding, "Estimated_Transaction_Cost_EUR")
    net_target = n(funding, "Net_Proceeds_EUR")
    add_table(
        document,
        ["Elemento Base", "Ipotesi"],
        [
            ["Target lordo cumulativo 2026–29", compact_euro(gross_target)],
            ["Costi transazione/campagna stimati", compact_euro(transaction_cost) + " · 5%"],
            ["Proventi netti cumulativi", compact_euro(net_target)],
            ["Tranche iniziale Q4 2026", compact_euro(n(funding, "Round_1_Net_EUR"))],
            ["Follow-on condizionale Q2 2028", compact_euro(n(funding, "Round_2_Net_EUR"))],
            ["Equity cumulativa illustrativa", f"{float(funding['Illustrative_Cumulative_Equity_Low']):.0%}–{float(funding['Illustrative_Cumulative_Equity_High']):.0%}"],
        ],
        [8.5, 8.1],
    )
    add_body(document, "€5,6m non è una singola campagna iniziale: il modello prevede €1,5m netti di pre-seed e €4,1m di follow-on solo dopo milestone. Valutazione ed equity sono intervalli illustrativi, non termini approvati.")
    document.add_heading("Readiness crowdfunding", level=2)
    add_bullets(document, [
        "società, cap table, IP, conti e governance verificati;",
        "provider autorizzato, consulenza legale/fiscale e documentazione d'offerta;",
        "metodo di valutazione, diritti, diluizione e data room approvati;",
        "anchor interest, audience reale e piano post-campagna;",
        "scenario di fallimento o ritardo senza spendere capitale non disponibile.",
    ])
    add_body(document, "Il Regolamento (UE) 2020/1503 disciplina i servizi coperti e include una soglia di €5m su 12 mesi. Applicazione, struttura della raccolta e comunicazioni richiedono conferma aggiornata da counsel italiano e piattaforma autorizzata.")
    page_break(document)

    # 12 — Financials
    add_section_label(document, "11", "Proiezioni finanziarie", "Scenario Base · stima gestionale, non forecast")
    year_rows = []
    for year in ("2027", "2028", "2029"):
        year_rows.append([
            year,
            compact_euro(sum_year(base, "Total_Operating_Revenue_EUR", year)),
            compact_euro(sum_year(base, "Total_Operating_Costs_EUR", year)),
            compact_euro(sum_year(base, "Operating_Cash_Flow_EUR", year)),
            f"{int(sum(n(row, 'B2B_B2G_Active_Contracts') for row in base if row['Quarter'] == f'Q4 {year}'))}",
        ])
    add_table(document, ["Anno", "Ricavi operativi", "Costi operativi", "Cash flow operativo", "Contratti attivi Q4"], year_rows, [2.0, 4.0, 4.0, 4.0, 2.6])
    add_picture(document, finance_chart, 15.8, "Ricavi, costi e cassa provengono dal modello finanziario trimestrale verificato.")
    add_bullets(document, [
        f"Cassa finale Q4 2029: {compact_euro(n(base[-1], 'Closing_Cash_EUR'))}.",
        f"Runway indicativa Q4 2029: {n(base[-1], 'Runway_Months'):.1f} mesi.",
        "Finanziamenti e ricavi sono colonne separate; i grant restano distinguibili.",
        "Founder compensation inclusa nel payroll; nessun lavoro gratuito indefinito.",
    ])
    add_body(document, "Le traiettorie dipendono da MAU, conversione, contratti, costi provider, hiring e timing raccolta non validati. Il workbook contiene formule, fonti, assunzioni e check.")
    page_break(document)

    # 13 — Use of funds and risks
    add_section_label(document, "12", "Uso dei fondi e rischi")
    add_picture(document, use_funds_chart, 16.2)
    add_table(
        document,
        ["Rischio prioritario", "Mitigazione / gate"],
        [
            ["Domanda ripetuta debole", "Cohort e pilot prima della scala"],
            ["Dati scarsi, falsi o obsoleti", "Provenienza, unknown coverage, expiry e audit"],
            ["Provider costoso o non fedele", "PoC, quote, contratti portabili e notifica degrado"],
            ["Gap tecnico e founder dependency", "Senior hire, advisor, governance e ownership"],
            ["Privacy, licenze e liability", "Minimizzazione, registri fonte, claims review, Article 35 assessment"],
            ["Crowdfunding fallisce", "Scope ridotto, alternative e nessuna spesa su capitale non raccolto"],
        ],
        [6.0, 10.6],
    )
    add_body(document, "La riserva commerciale non è un budget da spendere automaticamente. Ogni incremento in paid acquisition, automotive, seconda piattaforma o copertura geografica richiede evidenza e autorizzazione.")
    page_break(document)

    # 14 — Investment proposition and disclaimer
    add_section_label(document, "13", "Proposta di investimento", "Capitale milestone-based per quattro prove")
    add_bullets(document, [
        "Caso d'uso concreto e prototipo open source già funzionante.",
        "Opportunità di piattaforma più ampia, con AI confinata e routing spiegabile.",
        "Percorsi di ricavo consumer, fleet, pubblico e API, ancora da validare.",
        "Community come rete di qualità solo con moderazione e provenienza misurabili.",
        "Espansione regionale e automotive subordinata a gate, non promessa.",
    ])
    add_callout(document, "Investire ora significa finanziare evidenza: comportamento, dati, delivery ed economia.", THEME["paper"])
    document.add_heading("Disclaimer", level=2)
    add_body(document, "Questo documento è una bozza strategica di lavoro, non un documento di offerta pubblica né consulenza finanziaria. Le proiezioni sono stime basate su assunzioni e i risultati effettivi possono differire materialmente. Equity crowdfunding e ogni investimento richiedono preparazione legale, regolamentare, societaria, fiscale e finanziaria; investire comporta rischio. Marchi, naming, domini, società, profilo fondatore e termini economici restano soggetti a verifica e approvazione.")
    document.add_heading("Stato delle verifiche", level=2)
    add_body(document, "Il codice e i documenti tecnici supportano il perimetro della beta. Mercato, disponibilità a pagare, dati nazionali, provider, costi, team, autorizzazioni automotive, crowdfunding, valutazione e biografia non sono ancora verificati per pubblicazione.")
    document.add_heading("Contatti", level=2)
    add_body(document, "Founder: Domenico Campanella Scali — contatto professionale e fotografia da aggiungere dopo approvazione. Repository: github.com/campsh98-creator/Lastrico")
    p = document.add_paragraph("\nMOTIA\nRoutes shaped around you.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = rgb(THEME["teal"])

    document.save(DOCX_PATH)
    print(f"Created {DOCX_PATH}")


if __name__ == "__main__":
    build()
